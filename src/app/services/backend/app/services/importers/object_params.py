"""Импорт демо-датасета объектов (xlsx организатора) в справочники.

Каждая строка листа превращается в запись `parameters`: единица измерения,
значение по умолчанию, диапазон и примечание-источник. Из этих записей
фронтенд собирает форму, а бэкенд валидирует ввод — поэтому импорт
идемпотентен (повторный запуск обновляет записи, а не плодит дубли).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import DataSource, ObjectType, Parameter, ParameterGroup
from app.seed.param_codes import (
    ENUM_OPTIONS,
    HINT_ENUM,
    code_meta,
    group_code,
    resolve_code,
)
from app.services.importers.normalize import clean_text, parse_bool, parse_number

# Лист → (код типа объекта, название, описание)
SHEET_MAP: dict[str, tuple[str, str, str]] = {
    "Склад": (
        "warehouse",
        "Склад",
        "Складской комплекс: хранение, обработка и отгрузка грузов. "
        "Базовые процессы — приёмка, хранение, отбор, комплектация и отгрузка.",
    ),
    "Аэропорт": (
        "airport",
        "Аэропорт",
        "Пассажирский и грузовой аэропорт: наземное обслуживание, обработка багажа, "
        "внутритерминальная логистика и уборка.",
    ),
    "Медучреждение": (
        "medical",
        "Медицинское учреждение",
        "Больница, поликлиника или диагностический центр: доставка питания, белья, "
        "медикаментов, биоматериалов и вывоз отходов.",
    ),
}

# Источник, который проставляем всем параметрам датасета.
DATASET_SOURCE = "Демо-датасет организатора (Датасеты_хакатон.xlsx)"


@dataclass
class ImportResult:
    object_types: int = 0
    groups: int = 0
    parameters: int = 0
    updated: int = 0
    unresolved: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            "object_types": self.object_types,
            "groups": self.groups,
            "parameters": self.parameters,
            "updated": self.updated,
            "unresolved": self.unresolved,
            "warnings": self.warnings,
        }


def _get_or_create_source(db: Session) -> DataSource:
    source = db.scalar(select(DataSource).where(DataSource.name == DATASET_SOURCE))
    if source is None:
        source = DataSource(
            name=DATASET_SOURCE,
            kind="organizer",
            is_verified=True,
            note="Файл предоставлен организатором для конкурсного набора данных.",
        )
        db.add(source)
        db.flush()
    return source


def _row_values(row: tuple) -> dict[str, object]:
    return {
        "name": clean_text(row[0]),
        "unit": clean_text(row[1]),
        "default": row[2],
        "min": row[3],
        "max": row[4],
        "note": clean_text(row[5]),
    }


def import_object_parameters(db: Session, path: Path) -> ImportResult:
    """Загружает листы Склад / Аэропорт / Медучреждение в object_types + parameters."""
    result = ImportResult()
    source = _get_or_create_source(db)
    workbook = openpyxl.load_workbook(path, data_only=True)

    for sheet_name, (code, title, description) in SHEET_MAP.items():
        if sheet_name not in workbook.sheetnames:
            result.warnings.append(f"Лист «{sheet_name}» не найден в файле {path.name}")
            continue

        object_type = db.scalar(select(ObjectType).where(ObjectType.code == code))
        if object_type is None:
            object_type = ObjectType(
                code=code,
                name=title,
                description=description,
                icon=code,
                order_index=list(SHEET_MAP).index(sheet_name),
            )
            db.add(object_type)
            db.flush()
            result.object_types += 1
        else:
            object_type.name = title
            object_type.description = description
            result.updated += 1

        groups: dict[str, ParameterGroup] = {
            g.code: g for g in object_type.groups
        }
        existing_params = {p.code: p for p in object_type.parameters}

        current_group: ParameterGroup | None = None
        order_in_group = 0
        order_global = 0

        for raw_row in workbook[sheet_name].iter_rows(min_row=3, values_only=True):
            values = _row_values(raw_row)
            if not values["name"]:
                continue
            name = str(values["name"])

            # Строка-разделитель секции: «▌ РЕЖИМ РАБОТЫ»
            if name.lstrip().startswith("▌"):
                gcode = group_code(name)
                if gcode not in groups:
                    groups[gcode] = ParameterGroup(
                        object_type_id=object_type.id,
                        code=gcode,
                        name=name.lstrip("▌ ").strip(),
                        order_index=len(groups),
                    )
                    db.add(groups[gcode])
                    result.groups += 1
                current_group = groups[gcode]
                order_in_group = 0
                continue

            param_code = resolve_code(name)
            if param_code is None:
                result.unresolved.append(f"{sheet_name}: {name}")
                continue

            hint, required = code_meta(param_code)
            unit = values["unit"]
            unit = None if unit in {"-", ""} else unit

            default_num: Decimal | None = None
            default_text: str | None = None
            if hint in {"number", "integer", "percent"}:
                default_num = parse_number(values["default"])
                if default_num is None:
                    default_text = clean_text(values["default"])
                    if default_text is not None:
                        hint = "text"
            else:
                default_text = clean_text(values["default"])
                if hint == HINT_ENUM and default_text:
                    hint = HINT_ENUM

            min_value = parse_number(values["min"])
            max_value = parse_number(values["max"])
            # Датасет задаёт min/max и для текстовых полей (например, «Да»);
            # такие границы в числовую модель не переносятся.
            if hint in {"text", HINT_ENUM}:
                min_value = max_value = None

            options = ENUM_OPTIONS.get(param_code) if hint == HINT_ENUM else None

            values_map: dict[str, object] = {
                "object_type_id": object_type.id,
                "group_id": current_group.id if current_group else None,
                "name": name,
                "unit": unit,
                "value_type": hint,
                "default_value": default_num,
                "default_text": default_text,
                "min_value": min_value,
                "max_value": max_value,
                "options": options,
                "required": required,
                "is_demo": True,
                "order_index": order_global,
                "note": values["note"],
                "data_source_id": source.id,
            }

            existing = existing_params.get(param_code)
            if existing is None:
                db.add(Parameter(code=param_code, **values_map))
                result.parameters += 1
            else:
                for key, value in values_map.items():
                    setattr(existing, key, value)
                result.updated += 1

            order_in_group += 1
            order_global += 1

        db.flush()

    return result
